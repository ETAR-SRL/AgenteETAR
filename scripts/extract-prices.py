"""Local-only extraction. Never evaluates formulas, macros, links or document instructions."""
import json, sys, os, warnings, zipfile
from pathlib import Path
warnings.filterwarnings('ignore')
MAX_ROWS, MAX_PAGES, MAX_CELLS = 2000, 20, 40

def group_cells(cells):
    groups=[]
    for x,y,h,text,score in sorted(cells,key=lambda c:(c[1],c[0])):
        group=next((g for g in reversed(groups[-3:]) if abs(g['y']-y)<max(5,h*.55)),None)
        if group is None: group={'y':y,'cells':[]};groups.append(group)
        group['cells'].append((x,text,score))
    return [[c[1] for c in sorted(g['cells'])] for g in groups], [min(c[2] for c in g['cells']) for g in groups]

def extract(path):
    ext=Path(path).suffix.lower(); tables=[]; notices=[]; engine=None; context=[]
    def add(name,rows,scores=None,ocr=False):
        if len(rows)>MAX_ROWS or any(len(r)>MAX_CELLS for r in rows): raise ValueError('Demasiadas filas o columnas. Dividí el archivo (máximo 2000 filas y 40 columnas).')
        tables.append({'name':name,'rows':rows,'scores':scores or [],'ocr':ocr})
    def ocr_image(image,name):
        nonlocal engine
        from PIL import ImageOps
        import numpy as np
        image=ImageOps.exif_transpose(image)
        if image.width*image.height>30_000_000: raise ValueError('La imagen supera 30 megapíxeles. Reducí su tamaño.')
        if engine is None:
            from rapidocr_onnxruntime import RapidOCR
            engine=RapidOCR(intra_op_num_threads=2,inter_op_num_threads=1)
        result,_=engine(np.array(image.convert('RGB')))
        cells=[]
        for box,text,score in result or []:
            cells.append((min(p[0] for p in box),sum(p[1] for p in box)/4,max(p[1] for p in box)-min(p[1] for p in box),text,float(score)))
        rows,scores=group_cells(cells);add(name,rows,scores,True)
    if ext=='.xlsx':
        import openpyxl
        with zipfile.ZipFile(path) as z:
            if len(z.infolist())>10000 or sum(i.file_size for i in z.infolist())>100_000_000: raise ValueError('Excel expandido demasiado grande.')
        book=openpyxl.load_workbook(path,read_only=True,data_only=True,keep_links=False)
        try:
            if len(book.worksheets)>20: raise ValueError('Máximo 20 hojas por archivo.')
            for sheet in book:
                if sheet.max_row>MAX_ROWS+50 or sheet.max_column>MAX_CELLS: raise ValueError('Hoja demasiado grande. Exportá sólo la lista de precios.')
                rows=[]
                for row in sheet.iter_rows():
                    values=[]
                    for cell in row:
                        v=cell.value
                        if isinstance(v,(int,float)) and not isinstance(v,bool) and cell.number_format and set(cell.number_format)=={'0'} and len(cell.number_format)>1:
                            v=str(int(v)).zfill(len(cell.number_format))
                        if v is None:v=''
                        if not isinstance(v,(str,int,float)):v=str(v)
                        values.append(v)
                    rows.append(values)
                add(sheet.title,rows)
        finally:book.close()
        notices.append('Excel: se leen valores guardados; no se ejecutan fórmulas, macros ni vínculos. Revisá que las fórmulas se hayan recalculado y guardado en Excel.')
    elif ext=='.xls':
        import xlrd
        book=xlrd.open_workbook(path,on_demand=True)
        try:
            if book.nsheets>20: raise ValueError('Máximo 20 hojas por archivo.')
            for sheet in book.sheets():
                if sheet.nrows>MAX_ROWS or sheet.ncols>MAX_CELLS:raise ValueError('Hoja demasiado grande.')
                rows=[]
                for i in range(sheet.nrows):
                    values=[int(v) if isinstance(v,float) and v.is_integer() else v for v in sheet.row_values(i)]
                    rows.append(values)
                add(sheet.name,rows)
        finally:book.release_resources()
        notices.append('XLS: valores guardados, sin ejecutar fórmulas. Revisá los códigos que comienzan con ceros.')
    elif ext=='.pdf':
        import pdfplumber
        import pypdfium2 as pdfium
        with pdfplumber.open(path) as pdf:
            if len(pdf.pages)>MAX_PAGES:raise ValueError('Máximo 20 páginas. Dividí la lista en varios archivos.')
            for i,page in enumerate(pdf.pages):
                found=page.extract_tables()
                if found:
                    context.append(page.extract_text() or '')
                    for j,rows in enumerate(found):add(f'Página {i+1}, tabla {j+1}',[[v or '' for v in r] for r in rows])
                elif (page.extract_text() or '').strip():
                    words=page.extract_words(x_tolerance=4,y_tolerance=3,keep_blank_chars=True)
                    rows,scores=group_cells([(w['x0'],(w['top']+w['bottom'])/2,w['bottom']-w['top'],w['text'],1) for w in words]);add(f'Página {i+1}',rows,scores)
                else:
                    doc=pdfium.PdfDocument(path)
                    try:
                        p=doc[i]
                        if p.get_width()*p.get_height()*4>30_000_000:raise ValueError('Página escaneada demasiado grande.')
                        bitmap=p.render(scale=2);ocr_image(bitmap.to_pil(),f'Página {i+1} · OCR');bitmap.close();p.close()
                    finally:doc.close()
    else:
        from PIL import Image
        Image.MAX_IMAGE_PIXELS=30_000_000
        with Image.open(path) as image:ocr_image(image,'Imagen · OCR')
    if sum(len(t['rows']) for t in tables)>MAX_ROWS:raise ValueError('Máximo 2000 filas por archivo.')
    if any(t['ocr'] for t in tables):notices.append('Lectura OCR: verificá especialmente códigos, separadores decimales y moneda contra el original.')
    output={'tables':tables,'notices':notices,'context':'\n'.join(context)}
    if len(json.dumps(output))>500_000:raise ValueError('El texto extraído es demasiado extenso. Dividí el archivo.')
    return output

if __name__=='__main__':
    try:
        print(json.dumps(extract(sys.argv[1]),ensure_ascii=True,allow_nan=False))
    except Exception as error:
        print(json.dumps({'error':str(error)[:240]},ensure_ascii=True));sys.exit(1)
